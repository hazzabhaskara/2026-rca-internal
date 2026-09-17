"""
Phase D: Generate DAMA_Aged_Care_Research_v2.xlsx
Output: DAMA folder -> DAMA_Aged_Care_Research_v2.xlsx
- Uses enriched data with State/Region from ABN Lookup
- 2 new columns: registered_state (ABN) + la_coverage (National/Regional)
- Multi-state companies shown with all states (e.g. NSW, VIC, WA)
"""

import pandas as pd
from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings, os
warnings.filterwarnings('ignore')

print("="*60)
print("PHASE D: Generating DAMA_Aged_Care_Research_v2.xlsx")
print("="*60)

# ============================================================
# LOAD ENRICHED DATA
# ============================================================
df = pd.read_csv(r'c:\laragon\www\Ujicoba_RembulanCitraAbadi\la_enriched_temp.csv')
acila_df = pd.read_csv(r'c:\laragon\www\Ujicoba_RembulanCitraAbadi\acila_occupations_temp.csv')
dama_occ_df = pd.read_csv(r'c:\laragon\www\Ujicoba_RembulanCitraAbadi\dama_occupations_temp.csv')
today_str = date.today().strftime('%d %B %Y')

# Fill NaN in key columns
df['state_region_final'] = df['state_region_final'].fillna('Unknown')
df['registered_state'] = df['registered_state'].fillna('')
df['la_coverage'] = df['la_coverage'].fillna('')
df['registered_postcodes'] = df['registered_postcodes'].fillna('')

print(f"Loaded {len(df)} records")

# ============================================================
# STATISTICS
# ============================================================
total = len(df)
total_active = len(df[df['status'] == 'Active'])
total_dama = len(df[df['la_category'] == 'DAMA'])
total_aged_care = len(df[df['is_aged_care_related'] == True])
total_acila = len(df[df['latype'] == 'Industry - Aged Care'])

sector_counts = df['sector'].value_counts().reset_index()
sector_counts.columns = ['Sector', 'Count']
sector_counts['% Total'] = (sector_counts['Count'] / total * 100).round(1).astype(str) + '%'

# DAMA aged care summary
dama_aged = df[df['is_aged_care_related'] == True].groupby('latype').size().reset_index(name='Count')
dama_aged.columns = ['DAMA / LA Type', 'Aged Care Companies']
dama_aged = dama_aged.sort_values('Aged Care Companies', ascending=False)

# State breakdown for aged care
ac_by_state = df[df['is_aged_care_related'] == True].copy()
ac_by_state['state_display'] = ac_by_state['state_region_final']
state_summary = ac_by_state.groupby('state_display').size().reset_index(name='Count')
state_summary.columns = ['State / Region', 'Aged Care Companies']
state_summary = state_summary.sort_values('Aged Care Companies', ascending=False)

# ============================================================
# PREPARE SHEETS
# ============================================================

# Columns for All Companies
all_companies = df[[
    'companyname', 'latype', 'la_category', 'sector',
    'state_region_final', 'la_coverage', 'registered_state', 'registered_postcodes',
    'startdate', 'enddate', 'status', 'is_aged_care_related'
]].copy()
all_companies.columns = [
    'Company Name', 'LA Type', 'LA Category', 'Sector',
    'State / Region', 'LA Coverage', 'Registered State (ABN)', 'Postcodes',
    'Start Date', 'End Date', 'Status', 'Aged Care Related'
]
all_companies = all_companies.sort_values(['LA Category', 'Sector', 'Company Name'])

# DAMA Companies
dama_companies = df[df['la_category'] == 'DAMA'][[
    'companyname', 'latype', 'sector', 'state_region_final', 'la_coverage',
    'startdate', 'enddate', 'status', 'is_aged_care_related'
]].copy()
dama_companies.columns = [
    'Company Name', 'DAMA Type', 'Sector', 'State / Region', 'LA Coverage',
    'Start Date', 'End Date', 'Status', 'Aged Care Related'
]
dama_companies = dama_companies.sort_values(['DAMA Type', 'Sector', 'Company Name'])

# Aged Care Companies
aged_care_df = df[df['is_aged_care_related'] == True][[
    'companyname', 'latype', 'la_category', 'sector',
    'state_region_final', 'la_coverage', 'registered_state', 'registered_postcodes',
    'startdate', 'enddate', 'status'
]].copy()
aged_care_df.columns = [
    'Company Name', 'LA Type', 'LA Category', 'Sector',
    'State / Region', 'LA Coverage', 'Registered State (ABN)', 'Postcodes',
    'Start Date', 'End Date', 'Status'
]
aged_care_df = aged_care_df.sort_values(['LA Category', 'State / Region', 'Company Name'])

# Company Specific
company_specific_df = df[df['la_category'] == 'Company Specific'][[
    'companyname', 'latype', 'sector',
    'state_region_final', 'registered_state', 'registered_postcodes',
    'startdate', 'enddate', 'status', 'is_aged_care_related'
]].copy()
company_specific_df.columns = [
    'Company Name', 'LA Type', 'Sector',
    'State / Region', 'Registered State (ABN)', 'Postcodes',
    'Start Date', 'End Date', 'Status', 'Aged Care Related'
]
company_specific_df = company_specific_df.sort_values(['State / Region', 'Company Name'])

# Multi-state companies
multi_state_df = df[df['registered_state'].str.contains(',', na=False)][[
    'companyname', 'latype', 'la_category', 'sector',
    'registered_state', 'registered_postcodes', 'startdate', 'enddate', 'status'
]].copy()
multi_state_df.columns = [
    'Company Name', 'LA Type', 'LA Category', 'Sector',
    'States (All)', 'Postcodes', 'Start Date', 'End Date', 'Status'
]
multi_state_df = multi_state_df.sort_values(['States (All)', 'Company Name'])

print(f"\nSheets prepared:")
print(f"  All Companies: {len(all_companies)} rows")
print(f"  DAMA Companies: {len(dama_companies)} rows")
print(f"  Aged Care Companies: {len(aged_care_df)} rows")
print(f"  Company Specific: {len(company_specific_df)} rows")
print(f"  Multi-state Companies: {len(multi_state_df)} rows")
print(f"  ACILA Occupations: {len(acila_df)} rows")
print(f"  DAMA Occ: {len(dama_occ_df)} rows")

# ============================================================
# STYLES
# ============================================================
DARK_BLUE    = "1F3864"
MED_BLUE     = "2E75B6"
LIGHT_BLUE   = "BDD7EE"
AGED_PURPLE  = "7030A0"
LIGHT_PURPLE = "E2CFEE"
DAMA_GREEN   = "375623"
LIGHT_GREEN  = "C6EFCE"
ORANGE       = "C55A11"
LIGHT_ORANGE = "FCE4D6"
WHITE        = "FFFFFF"
DARK_GRAY    = "404040"
LIGHT_GRAY   = "F2F2F2"
MULTI_TEAL   = "1F6B75"
LIGHT_TEAL   = "DDEBF7"
COMPANY_BROWN = "833C00"
LIGHT_BROWN  = "FCE4D6"

def mf(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def mfont(bold=False, color=WHITE, size=10, italic=False):
    return Font(bold=bold, color=color, size=size, name='Calibri', italic=italic)

def mborder():
    s = Side(style='thin', color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def malign(h='left', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def style_title(ws, merge, text, bg, row_height=28, font_size=13):
    ws.merge_cells(merge)
    c = ws[merge.split(':')[0]]
    c.value = text
    c.font = mfont(bold=True, color=WHITE, size=font_size)
    c.fill = mf(bg)
    c.alignment = malign('center', 'center')
    ws.row_dimensions[int(merge.split(':')[0][1:])].height = row_height

def style_header(ws, row, col_start, col_end, bg):
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = mfont(bold=True, color=WHITE, size=9)
        cell.fill = mf(bg)
        cell.alignment = malign('center', 'center', wrap=True)
        cell.border = mborder()

def style_data_rows(ws, start_row, end_row, ncols, even_bg, odd_bg=WHITE,
                    special_col=None, special_vals=None, special_bg=None):
    for r in range(start_row, end_row + 1):
        is_even = (r - start_row) % 2 == 0
        bg = even_bg if is_even else odd_bg
        # Check special coloring
        if special_col and special_vals and special_bg:
            val = str(ws.cell(row=r, column=special_col).value or '')
            for sval, sbg in zip(special_vals, special_bg):
                if sval in val:
                    bg = sbg
                    break
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = mf(bg)
            cell.border = mborder()
            cell.font = mfont(bold=False, color=DARK_GRAY, size=9)
            cell.alignment = malign('left', 'center', wrap=True)

def auto_width(ws, min_w=12, max_w=48):
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, min_w), max_w)

def write_kpi(ws, col_letter, label, value, bg, row_start=4, row_end=7):
    ws.merge_cells(f'{col_letter}{row_start}:{col_letter}{row_end}')
    cell = ws[f'{col_letter}{row_start}']
    cell.value = f'{label}\n\n{value:,}' if isinstance(value, int) else f'{label}\n\n{value}'
    cell.font = Font(name='Calibri', size=11, bold=True, color=WHITE)
    cell.fill = mf(bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def write_table_block(ws, data_df, start_row, col_offset, header_bg, even_bg, odd_bg=WHITE):
    """Write df as mini table. Returns next free row."""
    cols = list(data_df.columns)
    # Header
    for j, col_name in enumerate(cols):
        c = ws.cell(row=start_row, column=col_offset + j, value=col_name)
        c.font = mfont(bold=True, color=WHITE, size=9)
        c.fill = mf(header_bg)
        c.alignment = malign('center', 'center', wrap=True)
        c.border = mborder()
    # Data
    for i, row in enumerate(data_df.itertuples(index=False), 1):
        bg = even_bg if i % 2 == 0 else odd_bg
        for j, val in enumerate(row):
            c = ws.cell(row=start_row + i, column=col_offset + j, value=val)
            c.fill = mf(bg)
            c.border = mborder()
            c.font = mfont(bold=False, color=DARK_GRAY, size=9)
            c.alignment = malign('left', 'center', wrap=False)
    return start_row + len(data_df) + 2

# ============================================================
# WRITE EXCEL
# ============================================================
output_path = r'c:\laragon\www\Ujicoba_RembulanCitraAbadi\DAMA\DAMA_Aged_Care_Research_v2.xlsx'
os.makedirs(os.path.dirname(output_path), exist_ok=True)

writer = pd.ExcelWriter(output_path, engine='openpyxl')

# Write data sheets
all_companies.to_excel(writer, sheet_name='All Companies', index=False, startrow=2)
dama_companies.to_excel(writer, sheet_name='DAMA Companies', index=False, startrow=2)
aged_care_df.to_excel(writer, sheet_name='Aged Care Companies', index=False, startrow=2)
company_specific_df.to_excel(writer, sheet_name='Company Specific', index=False, startrow=2)
multi_state_df.to_excel(writer, sheet_name='Multi-State Companies', index=False, startrow=2)
acila_df.to_excel(writer, sheet_name='ACILA Occupation List', index=False, startrow=2)
dama_occ_df.to_excel(writer, sheet_name='DAMA Aged Care Occupations', index=False, startrow=2)
pd.DataFrame().to_excel(writer, sheet_name='Summary Dashboard', index=False)

writer.close()

# ============================================================
# STYLE WITH OPENPYXL
# ============================================================
wb = load_workbook(output_path)

# ---- SUMMARY DASHBOARD ----
print("  Styling: Summary Dashboard...")
ws = wb['Summary Dashboard']
ws.sheet_properties.tabColor = DARK_BLUE

style_title(ws, 'A1:L1',
    'DAMA Labour Agreement Research  |  Aged Care Focus Dashboard',
    DARK_BLUE, row_height=36, font_size=16)

ws.merge_cells('A2:L2')
c = ws['A2']
c.value = f'Data current as at 30 April 2026  |  Generated: {today_str}  |  Source: immi.homeaffairs.gov.au + abr.business.gov.au'
c.font = mfont(bold=False, color="BBBBBB", size=9, italic=True)
c.fill = mf(DARK_GRAY)
c.alignment = malign('center', 'center')

ws.merge_cells('A3:L3')  # spacer

# KPI row
kpis = [
    ('A', 'Total Labour\nAgreements', total, DARK_BLUE),
    ('B', 'Active\nAgreements', total_active, "00B050"),
    ('C', 'DAMA\nCompanies', total_dama, MED_BLUE),
    ('D', 'Aged Care\nRelated', total_aged_care, AGED_PURPLE),
    ('E', 'Industry - Aged\nCare (ACILA)', total_acila, ORANGE),
    ('F', 'Multi-State\nCompanies', len(multi_state_df), MULTI_TEAL),
]
for col, lbl, val, bg in kpis:
    ws.column_dimensions[col].width = 18
    write_kpi(ws, col, lbl, val, bg, row_start=4, row_end=8)
    for r in range(4, 9):
        ws.row_dimensions[r].height = 18

# Sector table (col A, row 10)
ws.merge_cells('A10:D10')
hdr = ws['A10']
hdr.value = 'Labour Agreements by Sector'
hdr.font = mfont(bold=True, color=WHITE, size=11)
hdr.fill = mf(MED_BLUE)
hdr.alignment = malign('left', 'center')
ws.row_dimensions[10].height = 20
next_row = write_table_block(ws, sector_counts, 11, 1, MED_BLUE, LIGHT_BLUE)

# State breakdown table (col F, row 10)
ws.column_dimensions['F'].width = 30
ws.column_dimensions['G'].width = 20
ws.merge_cells('F10:G10')
hdr2 = ws['F10']
hdr2.value = 'Aged Care Companies by State'
hdr2.font = mfont(bold=True, color=WHITE, size=11)
hdr2.fill = mf(AGED_PURPLE)
hdr2.alignment = malign('left', 'center')
write_table_block(ws, state_summary, 11, 6, AGED_PURPLE, LIGHT_PURPLE)

# DAMA aged care table (col I, row 10)
ws.column_dimensions['I'].width = 35
ws.column_dimensions['J'].width = 16
ws.merge_cells('I10:J10')
hdr3 = ws['I10']
hdr3.value = 'Aged Care by DAMA / LA Type'
hdr3.font = mfont(bold=True, color=WHITE, size=11)
hdr3.fill = mf(DAMA_GREEN)
hdr3.alignment = malign('left', 'center')
write_table_block(ws, dama_aged, 11, 9, DAMA_GREEN, LIGHT_GREEN)

# ---- ALL COMPANIES ----
print("  Styling: All Companies...")
ws2 = wb['All Companies']
ws2.sheet_properties.tabColor = MED_BLUE
style_title(ws2, f'A1:{get_column_letter(len(all_companies.columns))}1',
    f'All Labour Agreement Companies ({total:,} records) — Enriched with State/Region',
    MED_BLUE)
style_header(ws2, 3, 1, len(all_companies.columns), MED_BLUE)

status_colors = {'Active': LIGHT_BLUE, 'Expiring Soon': LIGHT_ORANGE, 'Expired': 'FFCCCC'}
for r in range(4, 4 + len(all_companies)):
    status_val = str(ws2.cell(row=r, column=10).value or '')
    aged_val = str(ws2.cell(row=r, column=12).value or '')
    bg = LIGHT_GRAY if (r % 2 == 0) else WHITE
    if 'Expiring' in status_val: bg = LIGHT_ORANGE
    elif status_val == 'Expired': bg = 'FFCCCC'
    if aged_val.lower() in ('true', '1', 'yes'): bg = LIGHT_PURPLE
    for c in range(1, len(all_companies.columns) + 1):
        cell = ws2.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign()
auto_width(ws2)
ws2.freeze_panes = 'A4'

# ---- DAMA COMPANIES ----
print("  Styling: DAMA Companies...")
ws3 = wb['DAMA Companies']
ws3.sheet_properties.tabColor = DAMA_GREEN
style_title(ws3, f'A1:{get_column_letter(len(dama_companies.columns))}1',
    f'DAMA Companies ({total_dama:,} records) — 13 Regions', DAMA_GREEN)
style_header(ws3, 3, 1, len(dama_companies.columns), DAMA_GREEN)
for r in range(4, 4 + len(dama_companies)):
    aged_val = str(ws3.cell(row=r, column=9).value or '')
    bg = LIGHT_GREEN if r % 2 == 0 else WHITE
    if aged_val.lower() in ('true', '1', 'yes'): bg = LIGHT_PURPLE
    for c in range(1, len(dama_companies.columns) + 1):
        cell = ws3.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign()
auto_width(ws3)
ws3.freeze_panes = 'A4'

# ---- AGED CARE COMPANIES ----
print("  Styling: Aged Care Companies...")
ws4 = wb['Aged Care Companies']
ws4.sheet_properties.tabColor = AGED_PURPLE
style_title(ws4, f'A1:{get_column_letter(len(aged_care_df.columns))}1',
    f'Aged Care Companies ({total_aged_care}) — ACILA + DAMA + Company Specific', AGED_PURPLE)
style_header(ws4, 3, 1, len(aged_care_df.columns), AGED_PURPLE)

la_cat_colors = {'Industry': LIGHT_PURPLE, 'DAMA': LIGHT_GREEN,
                 'Company Specific': LIGHT_ORANGE, 'Skilled Refugee Pilot': LIGHT_BLUE}
for r in range(4, 4 + len(aged_care_df)):
    la_cat = str(ws4.cell(row=r, column=3).value or '')
    status_v = str(ws4.cell(row=r, column=11).value or '')
    bg = la_cat_colors.get(la_cat, LIGHT_GRAY)
    if status_v == 'Expiring Soon (<6mo)': bg = LIGHT_ORANGE
    elif status_v == 'Expired': bg = 'FFCCCC'
    for c in range(1, len(aged_care_df.columns) + 1):
        cell = ws4.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign(wrap=True)
auto_width(ws4)
ws4.freeze_panes = 'A4'

# ---- COMPANY SPECIFIC ----
print("  Styling: Company Specific...")
ws5 = wb['Company Specific']
ws5.sheet_properties.tabColor = COMPANY_BROWN
style_title(ws5, f'A1:{get_column_letter(len(company_specific_df.columns))}1',
    f'Company Specific Labour Agreements ({len(company_specific_df)} records)', COMPANY_BROWN)
style_header(ws5, 3, 1, len(company_specific_df.columns), COMPANY_BROWN)
for r in range(4, 4 + len(company_specific_df)):
    aged_val = str(ws5.cell(row=r, column=10).value or '')
    bg = LIGHT_BROWN if r % 2 == 0 else WHITE
    if aged_val.lower() in ('true', '1', 'yes'): bg = LIGHT_PURPLE
    for c in range(1, len(company_specific_df.columns) + 1):
        cell = ws5.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign()
auto_width(ws5)
ws5.freeze_panes = 'A4'

# ---- MULTI-STATE COMPANIES ----
print("  Styling: Multi-State Companies...")
ws6 = wb['Multi-State Companies']
ws6.sheet_properties.tabColor = MULTI_TEAL
style_title(ws6, f'A1:{get_column_letter(len(multi_state_df.columns))}1',
    f'Companies Operating in MULTIPLE States ({len(multi_state_df)} companies)', MULTI_TEAL)
style_header(ws6, 3, 1, len(multi_state_df.columns), MULTI_TEAL)
for r in range(4, 4 + len(multi_state_df)):
    states_val = str(ws6.cell(row=r, column=5).value or '')
    num_states = states_val.count(',') + 1
    if num_states >= 4: bg = 'FFF2CC'
    elif num_states >= 3: bg = LIGHT_TEAL
    else: bg = LIGHT_BLUE if r % 2 == 0 else WHITE
    for c in range(1, len(multi_state_df.columns) + 1):
        cell = ws6.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign()
auto_width(ws6)
ws6.freeze_panes = 'A4'

# ---- ACILA OCCUPATION LIST ----
print("  Styling: ACILA Occupation List...")
ws7 = wb['ACILA Occupation List']
ws7.sheet_properties.tabColor = ORANGE
style_title(ws7, f'A1:{get_column_letter(len(acila_df.columns))}1',
    'Aged Care Industry Labour Agreement (ACILA) — Eligible Occupations', ORANGE)
style_header(ws7, 3, 1, len(acila_df.columns), ORANGE)
skill_bgs = {'Skill Level 1': LIGHT_BLUE, 'Skill Level 2': LIGHT_GREEN, 'Skill Level 3': LIGHT_PURPLE}
for r in range(4, 4 + len(acila_df)):
    skill = str(ws7.cell(row=r, column=3).value or '')
    bg = skill_bgs.get(skill, LIGHT_GRAY)
    for c in range(1, len(acila_df.columns) + 1):
        cell = ws7.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign(wrap=True)
auto_width(ws7)

# ---- DAMA AGED CARE OCCUPATIONS ----
print("  Styling: DAMA Aged Care Occupations...")
ws8 = wb['DAMA Aged Care Occupations']
ws8.sheet_properties.tabColor = DAMA_GREEN
style_title(ws8, f'A1:{get_column_letter(len(dama_occ_df.columns))}1',
    'DAMA Regions — Aged Care Occupations Available per Region', DAMA_GREEN)
style_header(ws8, 3, 1, len(dama_occ_df.columns), DAMA_GREEN)
dama_row_colors = {
    'Northern Territory DAMA': 'D9F0A3',
    'South Australia Regional DAMA': 'C6EFCE',
    'Far North Queensland DAMA': 'BDD7EE',
    'WA DAMA': 'DDEBF7',
    'South West DAMA': 'EAF4FB',
    'Orana DAMA': 'FFF2CC',
    'Townsville DAMA': 'FCE4D6',
    'Great South Coast DAMA': 'E2EFDA',
    'Goldfields DAMA': 'FFF2CC',
    'East Kimberley DAMA': 'EDEDED',
    'Goulburn Valley DAMA': 'F4E8F9',
}
for r in range(4, 4 + len(dama_occ_df)):
    dn = str(ws8.cell(row=r, column=1).value or '')
    bg = dama_row_colors.get(dn, LIGHT_GRAY)
    for c in range(1, len(dama_occ_df.columns) + 1):
        cell = ws8.cell(row=r, column=c)
        cell.fill = mf(bg)
        cell.border = mborder()
        cell.font = mfont(color=DARK_GRAY, size=9)
        cell.alignment = malign(wrap=True)
auto_width(ws8)
ws8.freeze_panes = 'A4'

# ============================================================
# REORDER SHEETS
# ============================================================
desired_order = [
    'Summary Dashboard', 'Aged Care Companies', 'DAMA Companies',
    'Company Specific', 'Multi-State Companies', 'All Companies',
    'ACILA Occupation List', 'DAMA Aged Care Occupations'
]
for name in desired_order:
    if name in wb.sheetnames:
        idx_current = wb.sheetnames.index(name)
        idx_target = desired_order.index(name)
        wb.move_sheet(name, offset=idx_target - idx_current)

# ============================================================
# SAVE
# ============================================================
wb.save(output_path)
print(f"\nExcel saved to: {output_path}")
print(f"\nFINAL SUMMARY:")
print(f"  Total companies: {total:,}")
print(f"  Active: {total_active:,}")
print(f"  DAMA: {total_dama:,}")
print(f"  Aged Care related: {total_aged_care:,}")
print(f"  Multi-state companies: {len(multi_state_df):,}")
print(f"\n  Sheets: {wb.sheetnames}")
print("\nPHASE D COMPLETE!")
