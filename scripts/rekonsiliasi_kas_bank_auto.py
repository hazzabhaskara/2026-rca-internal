import os
import subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime

MYSQL_BIN = r"C:\laragon\bin\mysql\mysql-8.4.3-winx64\bin\mysql.exe"
DB_NAME = "rembulancitraabadi_2026_db"

def run_sql(db, sql):
    cmd = [MYSQL_BIN, "-u", "root", db, "-e", sql]
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        raise RuntimeError(f"MySQL Error: {res.stderr}")
    return res.stdout

def parse_tsv(raw):
    lines = [l for l in raw.strip().split("\n") if l.strip()]
    if not lines:
        return []
    headers = [h.strip() for h in lines[0].split("\t")]
    rows = []
    for line in lines[1:]:
        vals = [v.strip() for v in line.split("\t")]
        while len(vals) < len(headers):
            vals.append("")
        rows.append(dict(zip(headers, vals[:len(headers)])))
    return rows

def get_gl_balances():
    sql = """
    SELECT 
        c.kode_akun, 
        c.nama_akun, 
        COALESCE(SUM(jed.debits), 0) as total_debit, 
        COALESCE(SUM(jed.credits), 0) as total_credit,
        COALESCE(SUM(jed.debits - jed.credits), 0) as saldo_gl
    FROM chart_of_accounts c
    LEFT JOIN journal_entry_details jed ON jed.kode_akun = c.kode_akun
    WHERE c.kode_akun LIKE '1101%' OR c.kode_akun LIKE '1102%'
    GROUP BY c.kode_akun, c.nama_akun
    ORDER BY c.kode_akun;
    """
    data = parse_tsv(run_sql(DB_NAME, sql))
    res = {}
    for r in data:
        res[r['kode_akun']] = {
            'nama_akun': r['nama_akun'],
            'total_debit': float(r['total_debit']),
            'total_credit': float(r['total_credit']),
            'saldo_gl': float(r['saldo_gl'])
        }
    return res

# Definisi 4 Kategori Rekening 2026 & Data Fisik Master
MASTER_ACCOUNTS_2026 = [
    # KATEGORI 1: KAS & BANK RCA - PUSAT BANDUNG
    {
        'kategori': '1. KAS & BANK RCA - PUSAT BANDUNG',
        'kode': '1102001',
        'nama': 'MANDIRI 1310062136215_OPERASIONAL',
        'unit': 'PT RCA Pusat',
        'pic': 'Finance Pusat',
        'saldo_fisik_baseline': 243208723.97,
        'catatan_penyesuaian': 'Match s/d Juni. Perlu tarik e-statement Mandiri Q3 (Jul-Sep 2026).'
    },
    {
        'kategori': '1. KAS & BANK RCA - PUSAT BANDUNG',
        'kode': '1102002',
        'nama': 'MANDIRI 1310000650012_QRIS ODS',
        'unit': 'RCA Pusat / ODS',
        'pic': 'Finance Pusat / Ancol',
        'saldo_fisik_baseline': 3854543.49,
        'catatan_penyesuaian': 'Match s/d Juni. Rekening migrasi ke BNI QRIS.'
    },
    {
        'kategori': '1. KAS & BANK RCA - PUSAT BANDUNG',
        'kode': '1102003',
        'nama': 'MANDIRI 1310062146214_QRIS AWA',
        'unit': 'RCA Pusat / AWA',
        'pic': 'Finance Pusat / Ancol',
        'saldo_fisik_baseline': 6028321.94,
        'catatan_penyesuaian': 'Match s/d Juni (Selisih Rp 50 pembulatan). Rekening migrasi ke BNI.'
    },
    {
        'kategori': '1. KAS & BANK RCA - PUSAT BANDUNG',
        'kode': '1102004',
        'nama': 'MANDIRI 1310000620015_QRIS JBL',
        'unit': 'RCA Pusat / JBL',
        'pic': 'Finance Pusat / Ancol',
        'saldo_fisik_baseline': 447408.57,
        'catatan_penyesuaian': 'Match s/d Juni. Rekening migrasi ke BNI QRIS.'
    },
    {
        'kategori': '1. KAS & BANK RCA - PUSAT BANDUNG',
        'kode': '1102005',
        'nama': 'MANDIRI 1310062156213_QRIS DUFAN',
        'unit': 'RCA Pusat / Dufan',
        'pic': 'Finance Pusat / Ancol',
        'saldo_fisik_baseline': 641649.22,
        'catatan_penyesuaian': 'Match s/d Juni. Rekening migrasi ke BNI QRIS.'
    },

    # KATEGORI 2: KAS & BANK NON-RCA - BANDUNG & CABANG
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1101001',
        'nama': 'KAS OPERASIONAL BANDUNG',
        'unit': 'Kantor Pusat Bandung',
        'pic': 'Tim Ops Bandung (Dewi)',
        'saldo_fisik_baseline': 72874335.00,
        'catatan_penyesuaian': 'Target Fisik: Rp 72.874.335. Perlu Jurnal: Deviden Dufan -Rp 500Jt, Dobel -Rp 30Jt, CDM +Rp 70Jt.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1101003',
        'nama': 'KAS FIRMAN',
        'unit': 'Unit Firman',
        'pic': 'Kang Firman',
        'saldo_fisik_baseline': 33040889.00,
        'catatan_penyesuaian': 'Jan-Jul Match (Rp 33.040.889). Menunggu Kang Firman input mutasi nota Ags-Sep ke Excel.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1101004',
        'nama': 'KAS RISKI',
        'unit': 'Unit Riski (Bali)',
        'pic': 'Riski (Bali)',
        'saldo_fisik_baseline': -2114844.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1101005',
        'nama': 'KAS PESONA - IMAN',
        'unit': 'Pesona Waterpark',
        'pic': 'Iman Juarsa',
        'saldo_fisik_baseline': 4550911.00,
        'catatan_penyesuaian': '100% MATCH Sempurna (429 Mutasi klop).'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1101012',
        'nama': 'KAS QRIS PESONA',
        'unit': 'Pesona Waterpark',
        'pic': 'Iman Juarsa',
        'saldo_fisik_baseline': 0.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1102008',
        'nama': 'MANDIRI 1320571799999_PESONA',
        'unit': 'Pesona Waterpark',
        'pic': 'Finance Pesona',
        'saldo_fisik_baseline': 56587811.60,
        'catatan_penyesuaian': 'Mutasi 2026 match 100%. Selisih Rp 131.500 murni bawaan saldo awal bunga/adm 2025.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1102010',
        'nama': 'BCA 2330635366_YERRY',
        'unit': 'Lembang Zoo',
        'pic': 'Pak Yerry',
        'saldo_fisik_baseline': 137782962.10,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1102012',
        'nama': 'BCA EVI',
        'unit': 'Kings Bandung',
        'pic': 'Bu Evi',
        'saldo_fisik_baseline': 4403939.71,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '2. KAS & BANK BUKAN RCA - PUSAT BANDUNG & UNIT',
        'kode': '1101006',
        'nama': 'KAS PESONA - GOWET',
        'unit': 'Unit Wahana Go Wet',
        'pic': 'Iman / Tim Go Wet',
        'saldo_fisik_baseline': 4995970.00,
        'catatan_penyesuaian': 'File: 1101006_KAS PESONA - GOWET.xlsx (Maret 2026). Saldo DB: Rp 100.396 vs Fisik: Rp 4.995.970 | Gap: -Rp 4.895.574.'
    },

    # KATEGORI 3: KAS & BANK RCA - ANCOL
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1101002',
        'nama': 'KAS OMSET',
        'unit': 'Hub Ancol',
        'pic': 'Amel / Kasir Ancol',
        'saldo_fisik_baseline': 240666000.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1101008',
        'nama': 'KAS AMEL',
        'unit': 'Hub Sentral Ancol',
        'pic': 'Amel (Hub Sentral)',
        'saldo_fisik_baseline': 20000000.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1101010',
        'nama': 'KAS OPERASIONAL ANCOL',
        'unit': 'Ancol',
        'pic': 'Amel / Tim Ancol',
        'saldo_fisik_baseline': 14355899.00,
        'catatan_penyesuaian': 'Match (Selisih Rp 767.797 ter-offset beban operasional harian).'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1101011',
        'nama': 'KAS SETOR RETAIL',
        'unit': 'Ancol',
        'pic': 'Amel',
        'saldo_fisik_baseline': -45768470.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102006',
        'nama': 'MANDIRI AMEL 1200014486646_TRANS QR',
        'unit': 'Ancol',
        'pic': 'Amel',
        'saldo_fisik_baseline': 417986.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102007',
        'nama': 'BNI AMEL 1907207884_KAS OPS ANCOL',
        'unit': 'Ancol',
        'pic': 'Amel',
        'saldo_fisik_baseline': 3339286.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102015',
        'nama': 'BRI 034601053377509_AMEL',
        'unit': 'Ancol',
        'pic': 'Amel',
        'saldo_fisik_baseline': 384256.00,
        'catatan_penyesuaian': 'Match (Selisih -Rp 273 pembulatan).'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102016',
        'nama': 'BNI 2010790365_OPERASIONAL',
        'unit': 'Ancol',
        'pic': 'Finance Ancol',
        'saldo_fisik_baseline': 59712.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102017',
        'nama': 'BNI 2009826625_TAMPUNGAN QRIS DUFAN',
        'unit': 'Dufan Ancol',
        'pic': 'Finance Ancol',
        'saldo_fisik_baseline': 38136831.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102018',
        'nama': 'BNI 2020529741_TAMPUNGAN QRIS ODS JBL AWA',
        'unit': 'Wahana Ancol',
        'pic': 'Finance Ancol',
        'saldo_fisik_baseline': 8708339.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102019',
        'nama': 'BNI 2015889555_NEWSPAPER BOOTH',
        'unit': 'Ancol',
        'pic': 'Finance Ancol',
        'saldo_fisik_baseline': 2934828.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },
    {
        'kategori': '3. KAS & BANK RCA - ANCOL (DUFAN, ODS, AWA, JBL & PBN)',
        'kode': '1102020',
        'nama': 'BTN REDI',
        'unit': 'Ancol',
        'pic': 'Redi',
        'saldo_fisik_baseline': 8254.00,
        'catatan_penyesuaian': '100% MATCH Sempurna.'
    },

    # KATEGORI 4: REKENING STATIS / SALDO AWAL SAJA
    {
        'kategori': '4. REKENING KAS & BANK TANPA FILE SUMBER (STATIS / SALDO AWAL SAJA)',
        'kode': '1101007',
        'nama': 'KAS TAMPUNGAN OMSET',
        'unit': 'Pusat',
        'pic': 'Finance Pusat',
        'saldo_fisik_baseline': None,
        'catatan_penyesuaian': 'Saldo statis awal tahun (Rp 16 Jt mengendap).'
    },
    {
        'kategori': '4. REKENING KAS & BANK TANPA FILE SUMBER (STATIS / SALDO AWAL SAJA)',
        'kode': '1101009',
        'nama': 'KAS RAHMAN',
        'unit': 'Rahman',
        'pic': 'Rahman',
        'saldo_fisik_baseline': None,
        'catatan_penyesuaian': 'Nihil mutasi di 2026.'
    },
    {
        'kategori': '4. REKENING KAS & BANK TANPA FILE SUMBER (STATIS / SALDO AWAL SAJA)',
        'kode': '1102009',
        'nama': 'MANDIRI 1320077379999_PESONA GIRO',
        'unit': 'Pesona',
        'pic': 'Finance Pesona',
        'saldo_fisik_baseline': None,
        'catatan_penyesuaian': 'Saldo giro mengendap awal tahun.'
    },
    {
        'kategori': '4. REKENING KAS & BANK TANPA FILE SUMBER (STATIS / SALDO AWAL SAJA)',
        'kode': '1102011',
        'nama': 'BRI 34601001707562_AMEL',
        'unit': 'Amel',
        'pic': 'Amel',
        'saldo_fisik_baseline': None,
        'catatan_penyesuaian': 'Rekening lama digantikan BRI 7509.'
    },
    {
        'kategori': '4. REKENING KAS & BANK TANPA FILE SUMBER (STATIS / SALDO AWAL SAJA)',
        'kode': '1102013',
        'nama': 'BANK JAGO RAHMAN',
        'unit': 'Rahman',
        'pic': 'Rahman',
        'saldo_fisik_baseline': None,
        'catatan_penyesuaian': 'Saldo awal statis.'
    },
    {
        'kategori': '4. REKENING KAS & BANK TANPA FILE SUMBER (STATIS / SALDO AWAL SAJA)',
        'kode': '1102014',
        'nama': 'OCBC 693818420186_AMEL',
        'unit': 'Amel',
        'pic': 'Amel',
        'saldo_fisik_baseline': None,
        'catatan_penyesuaian': 'Saldo awal mengendap Rp 330.199.'
    }
]

def run_reconciliation():
    gl_data = get_gl_balances()
    os.makedirs("outputs", exist_ok=True)
    out_file = r"outputs\Rekonsiliasi_Kas_Bank_Terbaru.xlsx"

    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = "RINGKASAN_EKSEKUTIF_2026"
    ws_sum.views.sheetView[0].showGridLines = True

    # Styling Palette
    NAVY = "1F4E79"
    BLUE_HDR = "2F5597"
    LIGHT_BLUE = "D9E1F2"
    ZEBRA = "F2F5F9"
    GREEN_BG = "E2EFDA"
    GREEN_FG = "375623"
    RED_BG = "FCE4D6"
    RED_FG = "C65911"
    YELLOW_BG = "FFF2CC"
    YELLOW_FG = "7F6000"
    GRAY_BG = "F2F2F2"
    GRAY_FG = "595959"

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=9, italic=True, color="D9E1F2")
    font_cat_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_tbl_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_cell = Font(name="Calibri", size=10)
    font_cell_bold = Font(name="Calibri", size=10, bold=True)
    font_kpi_val = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    font_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="595959")

    thin_border = Border(
        left=Side(style='thin', color="D3D3D3"),
        right=Side(style='thin', color="D3D3D3"),
        top=Side(style='thin', color="D3D3D3"),
        bottom=Side(style='thin', color="D3D3D3")
    )

    # Title Banner
    ws_sum.merge_cells("A1:I1")
    ws_sum["A1"] = "LAPORAN REKONSILIASI KAS & BANK 2026: WEB INTERNAL VS FILE FISIK EXCEL"
    ws_sum["A1"].font = font_title
    ws_sum["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 30

    now_str = datetime.datetime.now().strftime("%d-%m-%Y %H:%M")
    ws_sum.merge_cells("A2:I2")
    ws_sum["A2"] = f"Generated by Automation: {now_str} WIB | Web DB: {DB_NAME} (http://2026-rca-internal.test/) | Posisi: 2026"
    ws_sum["A2"].font = font_sub
    ws_sum["A2"].fill = PatternFill("solid", fgColor=NAVY)
    ws_sum["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 18

    # Process Accounts
    processed_list = []
    current_cat = None
    count_match = 0
    count_selisih = 0
    count_pending = 0

    status_kas_bdg = ""
    status_kas_firman = ""

    for acc in MASTER_ACCOUNTS_2026:
        kode = acc['kode']
        gl_info = gl_data.get(kode, {'saldo_gl': 0.0, 'total_debit': 0.0, 'total_credit': 0.0})
        saldo_db = gl_info['saldo_gl']
        saldo_fisik = acc['saldo_fisik_baseline']
        
        if saldo_fisik is not None:
            selisih = saldo_db - saldo_fisik
            if abs(selisih) < 1:
                status = "MATCH / BALANCE"
                count_match += 1
            elif abs(selisih) <= 1000000:
                status = "SELISIH KECIL / TIMING"
                count_match += 1
            else:
                status = "ADA SELISIH"
                count_selisih += 1
        else:
            selisih = None
            status = "STATIS / SALDO AWAL"
            count_pending += 1

        # Monitoring Khusus Kas Bandung
        if kode == '1101001':
            if abs(saldo_db - 72874335.00) < 1:
                status_kas_bdg = "SUDAH ADJUSTED (Match Rp 72.874.335)"
                status = "MATCH (SETELAH ADJUSTMENT)"
            else:
                status_kas_bdg = f"BELUM DI-ADJUST (Saldo DB Rp {saldo_db:,.0f} vs Target Fisik Rp 72.874.335 | Gap Rp {selisih:,.0f})"

        # Monitoring Khusus Kas Firman
        if kode == '1101003':
            if abs(selisih) < 1:
                status_kas_firman = "EXCEL SUDAH UPDATE (Match Sempurna)"
                status = "MATCH (EXCEL TERUPDATE)"
            else:
                status_kas_firman = f"MENUNGGU UPDATE EXCEL AGUSTUS-SEPTEMBER (Saldo DB Rp {saldo_db:,.0f} vs Excel Juli Rp 33.040.889 | Gap Rp {selisih:,.0f})"

        processed_list.append({
            'kategori': acc['kategori'],
            'kode': kode,
            'nama': acc['nama'],
            'unit': acc['unit'],
            'pic': acc['pic'],
            'saldo_db': saldo_db,
            'saldo_fisik': saldo_fisik,
            'selisih': selisih,
            'status': status,
            'catatan': acc['catatan_penyesuaian']
        })

    # KPI Cards (Row 4-5)
    kpis = [
        ("TOTAL REKENING 2026", f"{len(processed_list)} Akun", LIGHT_BLUE),
        ("MATCH / SELARAS", f"{count_match} Akun", GREEN_BG),
        ("ADA SELISIH MATERIAL", f"{count_selisih} Akun", RED_BG),
        ("SALDO AWAL STATIS", f"{count_pending} Akun", GRAY_BG)
    ]
    col_pairs = [(1, 2), (3, 4), (5, 6), (7, 9)]
    for idx, (lbl, val, col_fill) in enumerate(kpis):
        c_start, c_end = col_pairs[idx]
        ws_sum.merge_cells(start_row=4, start_column=c_start, end_row=4, end_column=c_end)
        cell_lbl = ws_sum.cell(row=4, column=c_start, value=lbl)
        cell_lbl.font = font_kpi_lbl
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.fill = PatternFill("solid", fgColor=col_fill)

        ws_sum.merge_cells(start_row=5, start_column=c_start, end_row=5, end_column=c_end)
        cell_val = ws_sum.cell(row=5, column=c_start, value=val)
        cell_val.font = font_kpi_val
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.fill = PatternFill("solid", fgColor="FFFFFF")

        for r in [4, 5]:
            for c in range(c_start, c_end + 1):
                ws_sum.cell(row=r, column=c).border = thin_border

    # Headers for Table (Row 7)
    headers = [
        "Kode Akun", "Nama Akun Kas / Bank", "Entitas / Unit", "PIC",
        "Saldo Web 2026 (DB)", "Saldo Fisik (Excel)", "Selisih (DB - Fisik)",
        "Status Rekonsiliasi", "Catatan Rekonsiliasi & Tindak Lanjut"
    ]
    ws_sum.row_dimensions[7].height = 26
    for col_idx, h in enumerate(headers, 1):
        c = ws_sum.cell(row=7, column=col_idx, value=h)
        c.font = font_tbl_hdr
        c.fill = PatternFill("solid", fgColor=BLUE_HDR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    # Render Table Data with Category Separators
    cur_row = 8
    no_urut = 1
    for item in processed_list:
        if item['kategori'] != current_cat:
            current_cat = item['kategori']
            ws_sum.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=9)
            cat_cell = ws_sum.cell(row=cur_row, column=1, value=f">>> {current_cat}")
            cat_cell.font = font_cat_hdr
            cat_cell.fill = PatternFill("solid", fgColor="333F48")
            cat_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws_sum.row_dimensions[cur_row].height = 22
            cur_row += 1

        ws_sum.row_dimensions[cur_row].height = 20
        row_bg = ZEBRA if no_urut % 2 == 0 else "FFFFFF"

        vals = [
            item['kode'],
            item['nama'],
            item['unit'],
            item['pic'],
            item['saldo_db'],
            item['saldo_fisik'] if item['saldo_fisik'] is not None else "-",
            item['selisih'] if item['selisih'] is not None else "-",
            item['status'],
            item['catatan']
        ]

        for c_idx, val in enumerate(vals, 1):
            cell = ws_sum.cell(row=cur_row, column=c_idx, value=val)
            cell.font = font_cell
            cell.fill = PatternFill("solid", fgColor=row_bg)
            cell.border = thin_border

            # Number formatting
            if c_idx in [5, 6, 7] and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [1, 3, 4, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status styling
            if c_idx == 8:
                if "MATCH" in str(val):
                    cell.fill = PatternFill("solid", fgColor=GREEN_BG)
                    cell.font = Font(name="Calibri", size=9, bold=True, color=GREEN_FG)
                elif "SELISIH KECIL" in str(val):
                    cell.fill = PatternFill("solid", fgColor=YELLOW_BG)
                    cell.font = Font(name="Calibri", size=9, bold=True, color=YELLOW_FG)
                elif "ADA SELISIH" in str(val):
                    cell.fill = PatternFill("solid", fgColor=RED_BG)
                    cell.font = Font(name="Calibri", size=9, bold=True, color=RED_FG)
                else:
                    cell.fill = PatternFill("solid", fgColor=GRAY_BG)
                    cell.font = Font(name="Calibri", size=9, color=GRAY_FG)

        cur_row += 1
        no_urut += 1

    # Column Widths
    col_widths = {
        1: 14, 2: 36, 3: 24, 4: 20, 5: 22, 6: 22, 7: 22, 8: 24, 9: 55
    }
    for col_idx, width in col_widths.items():
        ws_sum.column_dimensions[get_column_letter(col_idx)].width = width

    # SHEET 2: MONITORING PENYESUAIAN KHUSUS
    ws_adj = wb.create_sheet(title="MONITORING_ADJUSTMENT")
    ws_adj.views.sheetView[0].showGridLines = True
    ws_adj.merge_cells("A1:F1")
    ws_adj["A1"] = "STATUS REAL-TIME PENYESUAIAN KAS BANDUNG & KAS FIRMAN"
    ws_adj["A1"].font = font_title
    ws_adj["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws_adj["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_adj.row_dimensions[1].height = 28

    adj_headers = ["Kode Akun", "Nama Akun", "Status Terkini", "Target Penyelarasan", "Langkah Perbaikan"]
    ws_adj.row_dimensions[3].height = 24
    for c_idx, h in enumerate(adj_headers, 1):
        c = ws_adj.cell(row=3, column=c_idx, value=h)
        c.font = font_tbl_hdr
        c.fill = PatternFill("solid", fgColor=BLUE_HDR)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    adj_data = [
        [
            "1101001",
            "KAS OPERASIONAL BANDUNG",
            status_kas_bdg,
            "Saldo DB menjadi Rp 72.874.335",
            "Posting Jurnal Penyesuaian: Deviden Dufan (Kredit Rp 500Jt), Koreksi Dobel (Kredit Rp 30Jt), CDM Yerry (Debet Rp 70Jt)."
        ],
        [
            "1101003",
            "KAS FIRMAN",
            status_kas_firman,
            "File Excel Firman terisi mutasi Ags-Sep",
            "Minta Kang Firman menginput mutasi pengeluaran nota operasional Agustus (Rp 205Jt) dan penerimaan modal (2x Rp 50Jt) ke Excel."
        ]
    ]

    for r_idx, row in enumerate(adj_data, 4):
        ws_adj.row_dimensions[r_idx].height = 25
        for c_idx, val in enumerate(row, 1):
            cell = ws_adj.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = thin_border
            if c_idx in [1]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if c_idx == 3:
                if "SUDAH" in str(val) or "MATCH" in str(val):
                    cell.fill = PatternFill("solid", fgColor=GREEN_BG)
                    cell.font = Font(name="Calibri", size=10, bold=True, color=GREEN_FG)
                else:
                    cell.fill = PatternFill("solid", fgColor=RED_BG)
                    cell.font = Font(name="Calibri", size=10, bold=True, color=RED_FG)

    col_widths_adj = {1: 14, 2: 30, 3: 50, 4: 35, 5: 65}
    for col_idx, width in col_widths_adj.items():
        ws_adj.column_dimensions[get_column_letter(col_idx)].width = width

    try:
        wb.save(out_file)
        print(f"SUCCESS: Laporan Rekonsiliasi 2026 disimpan di: {out_file}")
    except PermissionError:
        fallback_file = r"outputs\Rekonsiliasi_Kas_Bank_Terbaru_Live.xlsx"
        wb.save(fallback_file)
        print(f"INFO: File utama sedang dibuka di Excel, laporan disimpan di fallback: {fallback_file}")
        out_file = fallback_file
    print("\n" + "="*80)
    print("STATUS REAL-TIME PENYESUAIAN KHUSUS (SETIAP RUN OTOMASI):")
    print(f"1. KAS OPS BANDUNG (1101001) : {status_kas_bdg}")
    print(f"2. KAS FIRMAN      (1101003) : {status_kas_firman}")
    print("="*80 + "\n")
    return len(processed_list), count_match, count_selisih, count_pending

if __name__ == "__main__":
    run_reconciliation()
